import io

from tests.helpers import auth_cookies

ISO = "IDG holds a current ISO 9001 certification valid through 2027, verified annually."


DRAFT_ANSWER = "IDG holds a current ISO 9001 certification, drafted and reviewed in the cockpit."


def _stub_worker():
    """A worker builder whose answer() reports a StubProvider fallback (stub=True), so any draft
    path fails closed. Shaped to stand in for BOTH generation.assemble_worker (kwargs) and
    router.build_tenant_worker (positional request/tenant)."""
    async def _build(*_a, **_kw):
        from app.modules.agentcore.models import AgentResult, Critique, Decision

        class _W:
            async def answer(self, rid, text, **__):
                return AgentResult(requirement_id=rid, answer="", critique=Critique(grounding=0.0),
                                   decision=Decision.human_review, citations=[], stub=True)
        return _W()
    return _build


def _drafting_worker(answer: str = DRAFT_ANSWER):
    """A worker whose answer() returns a real (non-stub) grounded draft, so the background /extract
    job persists it onto the tree question -- this is the reviewed cockpit answer the /answers +
    export read paths now build from."""
    async def _build(*_a, **_kw):
        from app.modules.agentcore.models import AgentResult, Critique, Decision

        class _W:
            async def answer(self, rid, text, **__):
                return AgentResult(requirement_id=rid, answer=answer, critique=Critique(grounding=0.9),
                                   decision=Decision.auto_accept, citations=["doc:iso"], stub=False)
        return _W()
    return _build


def _setup_pack_with_requirements(client, monkeypatch, tenant: str, *, answer: str = DRAFT_ANSWER):
    """Create a pack, ingest KB, and drive the background /extract job (eager) to populate its flat
    requirements + tree from a canned extraction AND draft the KB question into the tree, offline.
    Returns (cookies, pack). The /answers + exports then read the reviewed tree, no fresh LLM."""
    c = auth_cookies(tenant, "owner")
    client.post("/api/v1/knowledge/ingest", json={"doc_id": "iso", "text": ISO}, cookies=c)
    pack = client.post("/api/v1/rfp/packs", json={"agent_id": "w1", "title": "Tender"}, cookies=c).json()

    import json as _json
    payload = _json.dumps([
        {"id": "A1", "title": "ISO Certification", "classification": "KB",
         "requirement": "Describe your ISO 9001 certification."},
        {"id": "A2", "title": "Pricing", "classification": "MANUAL", "requirement": "Complete pricing schedule."},
    ])

    def fake_get_provider(**_kwargs):
        class P:
            def generate(self, prompt, *, system=None, temperature=0.2, prefill=None):
                from app.modules.ai.gateway import LLMResponse
                return LLMResponse(text=payload)
        return P()

    monkeypatch.setattr("app.modules.rfp.generation.get_provider", fake_get_provider)
    monkeypatch.setattr("app.modules.rfp.generation.assemble_worker", _drafting_worker(answer))
    client.post(f"/api/v1/rfp/packs/{pack['id']}/extract", cookies=c,
                files={"file": ("t.txt", io.BytesIO(b"tender text"), "text/plain")})
    return c, pack


def test_extract_enforces_upload_cap(client, monkeypatch):
    # A tender upload larger than the per-endpoint cap is refused with 413 (shrink the cap so the
    # test payload stays small and never trips the global request-size middleware first).
    monkeypatch.setattr("app.modules.rfp.router.MAX_UPLOAD_BYTES", 16)
    c = auth_cookies("extrcap1", "owner")
    pack = client.post("/api/v1/rfp/packs", json={"agent_id": "w1", "title": "T"}, cookies=c).json()
    r = client.post(f"/api/v1/rfp/packs/{pack['id']}/extract", cookies=c,
                    files={"file": ("t.txt", io.BytesIO(b"x" * 64), "text/plain")})
    assert r.status_code == 413


def test_export_requires_extracted_requirements_first(client):
    c = auth_cookies("exportco0", "owner")
    pack = client.post("/api/v1/rfp/packs", json={"agent_id": "w1", "title": "Tender"}, cookies=c).json()
    assert client.get(f"/api/v1/rfp/packs/{pack['id']}/export/xlsx", cookies=c).status_code == 404


def test_export_xlsx_has_grounding_column_and_status(client, monkeypatch):
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "exportco1")

    r = client.get(f"/api/v1/rfp/packs/{pack['id']}/export/xlsx", cookies=c)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "Grounding %" in header
    assert "Status" in header
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    # A2 is MANUAL -> not drafted by AI, note says so
    manual_row = next(r for r in rows if r[0] == "A2")
    assert "team" in manual_row[5].lower()


def test_get_answers_json_for_cockpit(client, monkeypatch):
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "exportco1b")

    r = client.get(f"/api/v1/rfp/packs/{pack['id']}/answers", cookies=c)
    assert r.status_code == 200
    answers = r.json()
    assert len(answers) == 2
    kb = next(a for a in answers if a["req_id"] == "A1")
    # The KB answer is the reviewed cockpit draft on the tree, not a fresh read-path LLM call.
    assert kb["classification"] == "KB" and kb["answer"] == DRAFT_ANSWER
    assert kb["citations"] == ["doc:iso"] and "confidence" in kb
    manual = next(a for a in answers if a["req_id"] == "A2")
    assert manual["answer"] == "" and "team" in manual["note"].lower()


def _fail_if_called_worker(*_a, **_kw):
    raise AssertionError("the read path must NOT build a worker / call the generator")


def test_answers_read_path_uses_tree_and_never_calls_generator(client, monkeypatch):
    # After FIX 1 the plain (non-force) /answers read is built entirely from the reviewed tree --
    # it must never build a worker or fire a fresh LLM draft, no matter how many times it is read.
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "cacheco1")
    monkeypatch.setattr("app.modules.rfp.router.build_tenant_worker", _fail_if_called_worker)

    first = client.get(f"/api/v1/rfp/packs/{pack['id']}/answers", cookies=c)
    assert first.status_code == 200
    second = client.get(f"/api/v1/rfp/packs/{pack['id']}/answers", cookies=c)
    assert second.status_code == 200 and second.json() == first.json()
    kb = next(a for a in first.json() if a["req_id"] == "A1")
    assert kb["answer"] == DRAFT_ANSWER  # served straight from the tree


def test_answers_reflect_user_draft_then_locked_precedence(client, monkeypatch):
    # Precedence on the read/export path: locked user_draft > user_draft (saved) > ai_answer.
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "precedenceco")
    pid = pack["id"]

    # Locate the A1 question in the tree.
    sections = client.get(f"/api/v1/rfp/packs/{pid}/sections", cookies=c).json()
    q = next(q for s in sections for q in s["questions"] if q["req_id"] == "A1")

    # 1) ai_answer only -> the AI draft shows.
    a = client.get(f"/api/v1/rfp/packs/{pid}/answers", cookies=c).json()
    assert next(x for x in a if x["req_id"] == "A1")["answer"] == DRAFT_ANSWER

    # 2) a saved (unlocked) user_draft overrides the ai_answer.
    client.put(f"/api/v1/rfp/packs/{pid}/questions/{q['id']}/draft", cookies=c,
               json={"user_draft": "Reviewer edited draft."})
    a = client.get(f"/api/v1/rfp/packs/{pid}/answers", cookies=c).json()
    row = next(x for x in a if x["req_id"] == "A1")
    assert row["answer"] == "Reviewer edited draft." and row["note"] == ""

    # 3) locking that user_draft keeps it as the source of truth and flags it human-approved.
    client.post(f"/api/v1/rfp/packs/{pid}/questions/{q['id']}/lock", cookies=c)
    a = client.get(f"/api/v1/rfp/packs/{pid}/answers", cookies=c).json()
    row = next(x for x in a if x["req_id"] == "A1")
    assert row["answer"] == "Reviewer edited draft." and "human-approved" in row["note"].lower()

    # The clean submission .docx carries the reviewed (locked user_draft) text, not the AI draft.
    from docx import Document
    clean = client.get(f"/api/v1/rfp/packs/{pid}/export/docx-clean", cookies=c)
    text = "\n".join(p.text for p in Document(io.BytesIO(clean.content)).paragraphs)
    assert "Reviewer edited draft." in text and DRAFT_ANSWER not in text


def test_force_regenerate_raises_503_on_stub(client, monkeypatch):
    # The explicit editor-triggered force=true path may call the generator, but a StubProvider
    # fallback must fail closed with 503 so echo text is never persisted or exported.
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "stubforce")
    pid = pack["id"]
    # Undrafted-worker style: worker.answer() reports stub=True on the regenerate path.
    monkeypatch.setattr("app.modules.rfp.router.build_tenant_worker", _stub_worker())

    r = client.get(f"/api/v1/rfp/packs/{pid}/answers", params={"force": "true"}, cookies=c)
    assert r.status_code == 503
    # The pre-existing reviewed draft on the tree is untouched (nothing overwritten by the stub).
    sections = client.get(f"/api/v1/rfp/packs/{pid}/sections", cookies=c).json()
    q = next(q for s in sections for q in s["questions"] if q["req_id"] == "A1")
    assert q["ai_answer"] == DRAFT_ANSWER


def test_force_regenerate_redrafts_the_tree(client, monkeypatch):
    # A working provider on force=true re-drafts the tree questions, and the read reflects it.
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "regenco")
    pid = pack["id"]
    monkeypatch.setattr("app.modules.rfp.router.build_tenant_worker",
                        _drafting_worker("Regenerated grounded answer."))

    r = client.get(f"/api/v1/rfp/packs/{pid}/answers", params={"force": "true"}, cookies=c)
    assert r.status_code == 200
    assert next(x for x in r.json() if x["req_id"] == "A1")["answer"] == "Regenerated grounded answer."


def test_feedback_approve_and_reject(client):
    c = auth_cookies("fbco1", "owner")
    pack = client.post("/api/v1/rfp/packs", json={"agent_id": "w1", "title": "Tender"}, cookies=c).json()

    r = client.post(f"/api/v1/rfp/packs/{pack['id']}/feedback", cookies=c,
                    json={"req_id": "B1", "decision": "Approve", "comment": "Human-approved answer text."})
    assert r.status_code == 200
    assert r.json() == {"req_id": "B1", "decision": "Approve", "saved": True}

    r = client.post(f"/api/v1/rfp/packs/{pack['id']}/feedback", cookies=c,
                    json={"req_id": "B1", "decision": "Reject"})
    assert r.status_code == 200 and r.json()["decision"] == "Reject"


def test_feedback_rejects_invalid_decision(client):
    c = auth_cookies("fbco2", "owner")
    pack = client.post("/api/v1/rfp/packs", json={"agent_id": "w1", "title": "Tender"}, cookies=c).json()
    r = client.post(f"/api/v1/rfp/packs/{pack['id']}/feedback", cookies=c,
                    json={"req_id": "B1", "decision": "Bogus"})
    assert r.status_code == 422


def test_feedback_requires_editor_role(client):
    # A viewer must not be able to plant "verified ground truth" into the tenant's knowledge
    # base -- writing feedback needs the same Editor-or-above bar as every other mutation.
    owner = auth_cookies("fbco3", "owner")
    pack = client.post("/api/v1/rfp/packs", json={"agent_id": "w1", "title": "Tender"}, cookies=owner).json()
    viewer = auth_cookies("fbco3", "viewer")
    r = client.post(f"/api/v1/rfp/packs/{pack['id']}/feedback", cookies=viewer,
                    json={"req_id": "B1", "decision": "Approve", "comment": "should be blocked"})
    assert r.status_code == 403


def test_feedback_404s_for_nonexistent_pack(client):
    c = auth_cookies("fbco4", "owner")
    r = client.post("/api/v1/rfp/packs/does-not-exist/feedback", cookies=c,
                    json={"req_id": "B1", "decision": "Approve", "comment": "text"})
    assert r.status_code == 404


def test_force_regenerate_answers_requires_editor_role(client, monkeypatch):
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "fbco5")

    viewer = auth_cookies("fbco5", "viewer")
    r = client.get(f"/api/v1/rfp/packs/{pack['id']}/answers", params={"force": "true"}, cookies=viewer)
    assert r.status_code == 403
    # a plain (non-forced) read stays available to a viewer
    r2 = client.get(f"/api/v1/rfp/packs/{pack['id']}/answers", cookies=viewer)
    assert r2.status_code == 200


def test_export_docx_clean_and_review(client, monkeypatch):
    c, pack = _setup_pack_with_requirements(client, monkeypatch, "exportco2")

    clean = client.get(f"/api/v1/rfp/packs/{pack['id']}/export/docx-clean", cookies=c)
    assert clean.status_code == 200
    assert clean.headers["content-type"].startswith("application/vnd.openxmlformats")

    from docx import Document
    doc = Document(io.BytesIO(clean.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ISO Certification" in text
    assert "[To be completed by the team" in text  # MANUAL section placeholder

    review = client.get(f"/api/v1/rfp/packs/{pack['id']}/export/docx-review", cookies=c)
    assert review.status_code == 200
    review_doc = Document(io.BytesIO(review.content))
    review_text = "\n".join(p.text for p in review_doc.paragraphs)
    assert "Requirement:" in review_text

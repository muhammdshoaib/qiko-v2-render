"""Finance vertical: parser (synthetic xlsx), analytics computation, grounded chat (stub LLM),
and the tenant-scoped router (auth + role + BOLA isolation). No real OpenAI/AWS/network is used --
the chat path is exercised with a deterministic StubProvider injected into the service.
"""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.modules.agents.service import Agent
from app.modules.ai.gateway import StubProvider
from app.modules.identity.tokens import decode_token
from app.modules.finance.analytics import compute_analytics
from app.modules.finance.parser import ParseError, parse_workbook
from app.modules.finance.service import FinanceService, FinanceUploadError

from tests.helpers import auth_cookies

# --- synthetic workbook fixture --------------------------------------------------------------

_EXPENSE_HEADER = ["Property Reference", "Property Description", "Unit Ref", "Unit Description",
                   "Heading", "Supname", "Nett", "Vat", "Gross", "Exp date"]
_INCOME_HEADER = ["Property Reference", "Unit Ref", "Unit Description", "Amount", "Income Receivable",
                  "Tran date"]


def _build_workbook() -> bytes:
    """Two properties: PROP-A (profitable, low cost) and PROP-B (loss-making, high arrears)."""
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)
    # PROP-A: modest expenses across two vendors/categories
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
                "Acme Maintenance", 100.0, 20.0, 120.0, "01/03/2024"])
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Cleaning",
                "Sparkle Clean", 50.0, 10.0, 60.0, "05/03/2024"])
    # PROP-B: heavy expenses (loss-making) from one dominant vendor
    exp.append(["B002", "9 Mill Road", "B002/1", "Flat 1, 9 Mill Road", "Repairs",
                "Acme Maintenance", 800.0, 160.0, 960.0, "10/03/2024"])
    exp.append(["B002", "9 Mill Road", "B002/2", "Flat 2, 9 Mill Road", "Repairs",
                "Acme Maintenance", 500.0, 100.0, 600.0, "15/03/2024"])

    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    inc.append(["A001", "A001/1", "Flat 1, 12 High Street", 1000.0, 0.0, "01/03/2024"])
    inc.append(["B002", "B002/1", "Flat 1, 9 Mill Road", 500.0, 250.0, "10/03/2024"])
    inc.append(["B002", "B002/2", "Flat 2, 9 Mill Road", 400.0, 100.0, "15/03/2024"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- parser ----------------------------------------------------------------------------------

def test_parser_extracts_income_and_expense_rows():
    expenses, incomes = parse_workbook(_build_workbook())
    assert len(expenses) == 4
    assert len(incomes) == 3
    a_exp = [e for e in expenses if e.property_ref == "A001"]
    assert round(sum(e.gross for e in a_exp), 2) == 180.0
    assert {e.vendor for e in a_exp} == {"Acme Maintenance", "Sparkle Clean"}
    b_inc = [i for i in incomes if i.property_ref == "B002"]
    assert round(sum(i.income_receivable for i in b_inc), 2) == 350.0


def test_parser_rejects_non_xlsx_bytes():
    with pytest.raises(ParseError):
        parse_workbook(b"this is definitely not a spreadsheet")


def test_parser_handles_empty_workbook():
    wb = Workbook()
    wb.active.title = "Random"
    buf = io.BytesIO()
    wb.save(buf)
    expenses, incomes = parse_workbook(buf.getvalue())
    assert expenses == [] and incomes == []


# --- analytics -------------------------------------------------------------------------------

def test_analytics_computes_per_property_and_summary():
    expenses, incomes = parse_workbook(_build_workbook())
    result = compute_analytics(expenses, incomes)

    assert result["status"] == "success"
    summary = result["summary"]
    assert summary["propertyCount"] == 2
    assert summary["vendorCount"] == 2  # Acme Maintenance + Sparkle Clean
    assert summary["totalIncome"] == 1900.0  # 1000 + 500 + 400
    assert summary["totalExpenses"] == 1740.0  # 180 + 1560
    assert summary["netProfit"] == 160.0
    assert summary["outstanding"] == 350.0
    assert summary["expenseBreakdown"] == {"net": 1450.0, "vat": 290.0}

    props = {p["name"]: p for p in result["properties"]}
    prop_a = props["12 High Street"]
    assert prop_a["income"] == 1000.0
    assert prop_a["expenses"] == 180.0
    assert prop_a["netProfit"] == 820.0
    assert prop_a["marginPercent"] == 82.0
    assert prop_a["profitabilityClass"] == "high_performing"
    # Ported formula: costToIncome 180/1000*100 = 18 -> riskScore round(18*2.5) = 45 -> high.
    assert prop_a["costToIncomeRatio"] == 18.0
    assert prop_a["riskScore"] == 45

    prop_b = props["9 Mill Road"]
    assert prop_b["netProfit"] == -660.0  # 900 income - 1560 expenses
    assert prop_b["profitabilityClass"] == "under_performing"
    # cost-to-income 1560/900*100 = 173.33 -> riskScore round(173.33*2.5)=433 -> high
    assert prop_b["riskLevel"] == "high"


def test_analytics_risk_score_formula_is_cost_to_income_based():
    expenses, incomes = parse_workbook(_build_workbook())
    result = compute_analytics(expenses, incomes)
    prop_b = next(p for p in result["properties"] if p["name"] == "9 Mill Road")
    expected_cti = round((1560.0 / 900.0) * 100 * 100) / 100
    assert prop_b["costToIncomeRatio"] == expected_cti
    assert prop_b["riskScore"] == round(expected_cti * 2.5)


def test_analytics_summary_cards_field_mapping():
    """Regression coverage for the full qiko-fields-tracking.pdf summary contract."""
    expenses, incomes = parse_workbook(_build_workbook())
    result = compute_analytics(expenses, incomes)
    s = result["summary"]

    assert s["propertyCount"] == 2
    assert s["vendorCount"] == 2
    assert s["totalIncome"] == 1900.0
    assert s["totalExpenses"] == 1740.0
    assert s["netProfit"] == 160.0
    assert s["outstanding"] == 350.0
    assert s["expenseBreakdown"] == {"net": 1450.0, "vat": 290.0}

    # totalIncome section
    assert s["totalIncomeUnitCount"] == 3
    assert s["totalIncomeAveragePerProperty"] == pytest.approx(950.0, rel=1e-2)
    assert s["totalIncomeTopProperty"] == "12 High Street"
    assert s["totalIncomeTopPropertyRef"] == "A001"
    assert s["totalIncomeTopPropertyIncome"] == 1000.0

    # totalExpenses section
    assert s["totalExpensesNet"] == 1450.0
    assert s["totalExpensesVat"] == 290.0
    assert s["totalExpensesPropertyCount"] == 2
    assert s["totalExpensesInvoiceCount"] == 4
    assert s["totalExpensesAveragePerProperty"] == pytest.approx(870.0, rel=1e-2)
    assert s["totalExpensesTopVendor"] == "Acme Maintenance"
    assert s["totalExpensesTopVendorAmount"] == 1680.0
    assert s["totalExpensesTopCategory"] == "Repairs"
    assert s["totalExpensesTopCategoryAmount"] == 1680.0

    # netProfit section
    assert s["netProfitMarginPercent"] == pytest.approx(8.42, rel=1e-2)
    assert s["netProfitAverageMargin"] == pytest.approx(4.34, rel=1e-2)
    assert s["netProfitTopProperty"] == "12 High Street"
    assert s["netProfitTopPropertyRef"] == "A001"
    assert s["netProfitTopPropertyProfit"] == 820.0
    assert s["netProfitLowestPerformanceProperty"] == "9 Mill Road"
    assert s["netProfitLowestPerformancePropertyRef"] == "B002"
    assert s["netProfitLowestPerformancePropertyProfit"] == -660.0
    assert s["netProfitCostToIncomeRatio"] == pytest.approx(91.58, rel=1e-2)
    assert s["netProfitPropertyCount"] == 2

    # outstanding section
    assert s["outstandingPropertyCount"] == 1
    assert s["outstandingPositiveTotal"] == 350.0
    assert s["outstandingNegativeTotal"] == 0.0
    assert s["outstandingTopProperty"] == "9 Mill Road"
    assert s["outstandingTopPropertyRef"] == "B002"
    assert s["outstandingTopPropertyOutstanding"] == 350.0
    assert s["outstandingAveragePerProperty"] == pytest.approx(350.0, rel=1e-2)
    assert s["outstandingAsPercentOfIncome"] == pytest.approx(18.42, rel=1e-2)
    assert s["outstandingIncomeReceivableCount"] == 2

    # propertyCount / portfolio section
    assert s["highPerforming"] == 1
    assert s["averagePerforming"] == 0
    assert s["underPerforming"] == 1
    assert s["highRisk"] == 2
    assert s["mediumRisk"] == 0
    assert s["lowRisk"] == 0
    assert s["totalUnits"] == 3
    assert s["totalInvoices"] == 4
    assert s["averageIncomePerProperty"] == pytest.approx(950.0, rel=1e-2)
    assert s["averageExpensePerProperty"] == pytest.approx(870.0, rel=1e-2)
    assert s["topProperty"] == "12 High Street"
    assert s["topPropertyRef"] == "A001"
    assert s["worstProperty"] == "9 Mill Road"
    assert s["worstPropertyRef"] == "B002"


def test_analytics_insights_and_trends_present():
    expenses, incomes = parse_workbook(_build_workbook())
    result = compute_analytics(expenses, incomes)

    insights = result["insights"]
    assert insights["propertyInsights"]["topPerformer"]["propertyName"] == "12 High Street"
    assert insights["propertyInsights"]["worstPerformer"]["propertyName"] == "9 Mill Road"

    sections = {t["section"] for t in result["trends"]}
    assert {"Property Performance", "Arrears & Collections"} <= sections
    arrears = next(t for t in result["trends"] if t["section"] == "Arrears & Collections")
    total_arrears = next(i for i in arrears["insights"] if i["title"] == "Total Arrears")
    assert total_arrears["metricValue"] == 350.0

    charts = result["charts"]
    assert charts["topVendorsBySpend"][0]["vendorName"] == "Acme Maintenance"


def test_analytics_empty_portfolio_is_safe():
    result = compute_analytics([], [])
    assert result["summary"]["propertyCount"] == 0
    assert result["properties"] == []
    assert len(result["trends"]) == 4


def _workbook_with_owner_expenditure() -> bytes:
    """PROP-A is normal; PROP-OWNER is flagged as owner expenditure and should be excluded."""
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
                "Acme Maintenance", 100.0, 20.0, 120.0, "01/03/2024"])
    exp.append(["OWN1", "owner expenditure property", "OWN1/1", "Owner Expenditure", "Maintenance",
                "Owner Vendor", 500.0, 100.0, 600.0, "02/03/2024"])

    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    inc.append(["A001", "A001/1", "Flat 1, 12 High Street", 1000.0, 0.0, "01/03/2024"])
    inc.append(["OWN1", "OWN1/1", "Owner Expenditure", 300.0, 50.0, "02/03/2024"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_analytics_excludes_owner_expenditure_rows():
    expenses, incomes = parse_workbook(_workbook_with_owner_expenditure())
    result = compute_analytics(expenses, incomes)
    s = result["summary"]

    assert s["propertyCount"] == 1
    assert s["totalIncome"] == 1000.0
    assert s["totalExpenses"] == 120.0
    assert s["outstanding"] == 0.0
    assert s["vendorCount"] == 1
    assert s["totalExpensesTopVendor"] == "Acme Maintenance"
    assert s["totalInvoices"] == 1
    assert len(result["properties"]) == 1
    assert result["properties"][0]["propertyRef"] == "A001"


def _two_month_workbook() -> bytes:
    """One property with income + expenses split across Feb and Mar 2024, so a month filter can be
    shown to narrow the totals to a single month."""
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
                "Acme Maintenance", 100.0, 20.0, 120.0, "05/02/2024"])  # Feb
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
                "Acme Maintenance", 200.0, 40.0, 240.0, "10/03/2024"])  # Mar

    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    inc.append(["A001", "A001/1", "Flat 1, 12 High Street", 1000.0, 0.0, "05/02/2024"])  # Feb
    inc.append(["A001", "A001/1", "Flat 1, 12 High Street", 500.0, 0.0, "10/03/2024"])  # Mar

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_service_analytics_month_filter_narrows_to_that_month():
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-m", _two_month_workbook(), filename="p.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    full = await svc.analytics("acme", "agent-m")
    assert full["summary"]["totalIncome"] == 1500.0
    assert full["summary"]["totalExpenses"] == 360.0

    march = await svc.analytics("acme", "agent-m", month="2024-03")
    assert march["summary"]["totalIncome"] == 500.0
    assert march["summary"]["totalExpenses"] == 240.0
    # The month is expanded to that month's inclusive ISO range.
    assert march["dateRange"]["from"] == "2024-03-01"
    assert march["dateRange"]["to"] == "2024-03-31"

    # An explicit from/to still wins over a supplied month.
    feb = await svc.analytics("acme", "agent-m", frm="2024-02-01", to="2024-02-29", month="2024-03")
    assert feb["summary"]["totalIncome"] == 1000.0


async def test_service_analytics_all_time_date_range_uses_file_declared_periods_not_row_dates():
    # EL-1024: with no month filter ("All time"), dateRange used to fall back to the earliest and
    # latest *row* date across every upload, so a YTD workbook with arrears back to 2005 made the
    # Financial Chat's "Period: ..." line (built from this field, see chat.py _summary_lines) claim
    # a 21-year range for a portfolio that only ever had February 2026 and December 2025 uploaded.
    # It must instead span the files' own declared periods, exactly like availableMonths.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
               "Acme Maintenance", 100.0, 20.0, 120.0, "12/06/2005"])
    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    buf = io.BytesIO()
    wb.save(buf)
    await svc.upload("acme", "agent-dr", buf.getvalue(),
                     filename="Touchstone_Income_and_Exp_YTD_December_2025.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    await svc.upload("acme", "agent-dr", _single_month_workbook("15/02/2026"),
                     filename="Touchstone_Income_and_Exp_February_2026.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    full = await svc.analytics("acme", "agent-dr")
    assert full["dateRange"]["from"] == "2025-12-01"
    assert full["dateRange"]["to"] == "2026-02-28"
    # The actual 2005 arrears row must still be counted in the totals -- only the reported *label*
    # changes, real financial figures are never silently dropped from the sums.
    assert full["summary"]["totalExpenses"] == 120.0

    # An explicit month filter is unaffected: it still narrows to exactly that month.
    narrowed = await svc.analytics("acme", "agent-dr", month="2026-02")
    assert narrowed["dateRange"]["from"] == "2026-02-01"
    assert narrowed["dateRange"]["to"] == "2026-02-28"


# --- service (with injected stub LLM) --------------------------------------------------------

def _single_month_workbook(income_date: str) -> bytes:
    """One property, one income row on the given DD/MM/YYYY date, no expenses. Used to build
    workbooks that each cover a single, distinct month regardless of when they are uploaded."""
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)

    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    inc.append(["A001", "A001/1", "Flat 1, 12 High Street", 1000.0, 0.0, income_date])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_service_list_files_reports_available_months_from_each_files_own_name():
    # EL-1014 (part 1): the month dropdown used to be built from each file's *upload* timestamp
    # (FinanceFile.uploaded_at), not the period its data actually covers. That was fixed by reading
    # every transaction date out of the workbook instead.
    #
    # EL-1021/EL-1024 (follow-up): that row-scan fix then surfaced a *different* problem. A real
    # "YTD" workbook can carry individual invoices going back years (arrears, historic maintenance
    # records) alongside its current reporting period, so scanning every row's date made
    # availableMonths span a decade or more for a single upload -- e.g. two files nominally
    # "February 2026" and "December 2025" produced a dropdown (and a chat "Period: ..." line)
    # spanning back to 2005, even though only those two periods were ever uploaded. availableMonths
    # must instead reflect the ONE period each uploaded workbook itself declares (its own file
    # name, which this client's exports consistently end in, e.g.
    # "Touchstone_Income_and_Exp_February_2026.xlsx" or "..._YTD_December_2025.xlsx"), independent
    # of both upload time and how far back its row-level data happens to reach.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("15/02/2026"),
                     filename="Touchstone_Income_and_Exp_February_2026.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    await svc.upload("acme", "agent-1", _single_month_workbook("20/12/2025"),
                     filename="Touchstone_Income_and_Exp_YTD_December_2025.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == ["2025-12", "2026-02"]


async def test_service_available_months_ignores_historic_rows_scattered_inside_a_ytd_file():
    # The exact EL-1021 scenario: a YTD workbook's own row data spans years (arrears/historic
    # invoices) even though the file itself is nominally for one period. Only the file's own
    # declared period should surface, not every year its rows happen to touch.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
               "Acme Maintenance", 100.0, 20.0, 120.0, "12/06/2005"])  # old arrears row
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
               "Acme Maintenance", 50.0, 10.0, 60.0, "20/12/2025"])  # the file's real period
    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    buf = io.BytesIO()
    wb.save(buf)

    await svc.upload("acme", "agent-1", buf.getvalue(),
                     filename="Touchstone_Income_and_Exp_YTD_December_2025.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == ["2025-12"]


async def test_service_available_months_does_not_match_a_month_name_mid_word():
    # "SummaryMay_2026.xlsx" must not be misread as May 2026: "May" is preceded by a letter
    # ("...aryMay"), not a real separator, so it is part of a different word ("Summary").
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("15/02/2026"), filename="SummaryMay_2026.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == []


async def test_service_available_months_skips_a_file_with_no_recognizable_period_in_its_name():
    # A file whose name carries no Month/Year (e.g. a day-rates reference sheet) contributes
    # nothing to the dropdown rather than guessing from its row data.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("31/13/2026"), filename="Day_Rates.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == []


async def test_service_available_months_matches_month_and_year_with_no_separator():
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("15/09/2026"), filename="September2026.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == ["2026-09"]


async def test_service_available_months_rejects_an_implausible_year_next_to_a_month_name():
    # A month-like substring next to an unrelated 4-digit number (a reference/serial code, not a
    # year) must not fabricate a bogus period like "0007-09" -- the same class of "wildly wrong
    # date range" bug this whole fix removes, just inverted.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("15/09/2026"), filename="September_0007.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == []


async def test_service_available_months_uses_the_rightmost_period_in_a_dual_dated_name():
    # Documents a known, accepted limitation: a filename naming two periods (e.g. a "compare this
    # month against last year" export) resolves to whichever Month/Year comes last in the name.
    # This client's actual naming convention only ever has one, but if it ever names a range or
    # comparison report, this is the current behavior rather than an ambiguity error.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("15/09/2026"),
                     filename="September_2025_vs_September_2026.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == ["2026-09"]


async def test_service_available_months_prefers_source_name_over_file_name():
    # A curated display title (source_name, from the upload's "title" field) takes precedence
    # over the raw stored file_name when both are present and disagree.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-1", _single_month_workbook("15/02/2026"), filename="Day_Rates.xlsx",
                     source_name="Touchstone_Income_and_Exp_February_2026.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    listing = await svc.list_files("acme", "agent-1")
    assert listing["availableMonths"] == ["2026-02"]


async def test_service_analytics_all_time_date_range_is_blank_when_no_file_has_a_declared_period():
    # If not one uploaded file's name carries a recognizable period, the reported range must be
    # blank rather than silently falling back to compute_analytics's row-scan value -- that
    # fallback is exactly the multi-year-spread bug EL-1021/EL-1024 removes, so resurrecting it
    # here (even only for this rare case) would defeat the fix.
    svc = FinanceService(provider=StubProvider(canned="ok"))
    await svc.upload("acme", "agent-nn", _single_month_workbook("12/06/2005"), filename="portfolio.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    full = await svc.analytics("acme", "agent-nn")
    assert full["dateRange"] == {"from": "", "to": "", "label": ""}
    # The row is still fully counted in the totals -- only the reported range is withheld.
    assert full["summary"]["totalIncome"] == 1000.0


async def test_service_upload_parse_store_and_analytics():
    svc = FinanceService(provider=StubProvider(canned="ok"))
    stored = await svc.upload("acme", "agent-1", _build_workbook(), filename="portfolio.xlsx",
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert stored.row_count == 7
    assert stored.namespace.startswith("finance-agent-1-")

    listing = await svc.list_files("acme", "agent-1")
    assert listing["totalFiles"] == 1
    assert listing["totalRows"] == 7

    analytics = await svc.analytics("acme", "agent-1")
    assert analytics["summary"]["propertyCount"] == 2
    assert analytics["totalFiles"] == 1


async def test_service_rejects_oversize_and_bad_type():
    svc = FinanceService(provider=StubProvider(canned="ok"))
    with pytest.raises(FinanceUploadError):
        await svc.upload("acme", "agent-1", b"x" * 10, filename="f.xlsx", content_type="text/plain")
    with pytest.raises(FinanceUploadError):
        await svc.upload("acme", "agent-1", b"not-xlsx", filename="f.xlsx",
                         content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


async def test_service_chat_is_grounded_in_analytics_via_stub():
    # StubProvider echoes the CONTEXT block, proving the chat feeds real computed figures to the LLM.
    svc = FinanceService(provider=StubProvider())
    await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result = await svc.chat("acme", "agent-1", "What is the total income across the portfolio?")
    assert result["intent"] == "data"
    assert result["rowsRetrieved"] >= 1
    assert "1900" in result["reply"]  # the real total income appears in the grounded context

    history = await svc.history("acme", "agent-1")
    assert history["count"] == 2  # user + assistant
    assert history["history"][0]["role"] == "user"


async def test_service_chat_follow_up_sees_the_prior_exchange():
    # EL-1003: "why do you say that" must resolve against what was actually asked/answered, not
    # just be persisted for the transcript while the model itself gets no memory of it.
    svc = FinanceService(provider=StubProvider())
    await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    await svc.chat("acme", "agent-1", "Which property has the highest risk score?")
    follow_up = await svc.chat("acme", "agent-1", "why do you say that")
    assert "RECENT CONVERSATION" in follow_up["reply"]
    assert "highest risk score" in follow_up["reply"]


async def test_service_chat_includes_vendor_names_for_vendor_questions():
    # EL-1007: vendor questions used to get no vendor names in CONTEXT at all, so the model
    # (correctly, per its own system prompt) refused rather than invent one.
    svc = FinanceService(provider=StubProvider())
    await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result = await svc.chat("acme", "agent-1", "Who is my biggest vendor by spend?")
    assert "VENDORS" in result["reply"]
    assert "Acme Maintenance" in result["reply"]


async def test_service_chat_includes_unit_level_data_for_unit_questions():
    # EL-1008: unit-level questions used to get nothing finer-grained than the property row.
    svc = FinanceService(provider=StubProvider())
    await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result = await svc.chat("acme", "agent-1", "Show me the unit level breakdown for 12 High Street")
    assert "UNITS" in result["reply"]
    assert "A001/1" in result["reply"]


def test_normalize_markdown_tables_leaves_well_formed_input_untouched():
    from app.modules.finance.chat import _normalize_markdown_tables

    already_fine = (
        "Some prose that mentions true || false without a real table.\n\n"
        "- a bullet\n- another bullet\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |"
    )
    assert _normalize_markdown_tables(already_fine) == already_fine


def test_normalize_markdown_tables_does_not_drop_a_cell_that_looks_like_a_separator():
    # A "--"/"-" no-data placeholder is common in real financial tables and, in isolation, matches
    # the same dashes pattern as a genuine header-separator row. Guessing the column width from it
    # must not silently delete the other cells in that ordinary, already-correct row.
    from app.modules.finance.chat import _normalize_markdown_tables

    row_with_placeholder_cell = "| A | B | -- |"
    assert _normalize_markdown_tables(row_with_placeholder_cell) == row_with_placeholder_cell


def test_normalize_markdown_tables_reflows_a_glued_table_that_has_a_placeholder_cell():
    # A "--" no-data placeholder inside one of several genuinely glued rows must not stop the
    # whole table (which has a real, unambiguous separator row) from being reflowed correctly.
    from app.modules.finance.chat import _normalize_markdown_tables

    glued = "| Property | Arrears | |------|---------|| A | -- || B | 120.00 |"
    assert _normalize_markdown_tables(glued) == (
        "| Property | Arrears |\n|------|---------|\n| A | -- |\n| B | 120.00 |"
    )


def test_normalize_markdown_tables_handles_prose_before_the_table_on_the_same_line():
    from app.modules.finance.chat import _normalize_markdown_tables

    glued = "Here's the data: | A | B ||---|---|| 1 | 2 |"
    assert _normalize_markdown_tables(glued) == (
        "Here's the data:\n| A | B |\n|---|---|\n| 1 | 2 |"
    )


def test_normalize_markdown_tables_documents_ragged_empty_cell_limitation():
    # Known residual limitation (see _reflow_glued_table_row's docstring): a ragged glued table
    # whose extra/missing cells are empty, rather than containing visible text, can pass both
    # safety checks and get reflowed with values reattributed to the wrong column. Values are
    # preserved (nothing is deleted), but "9" and "7" end up under different columns than the
    # input implied. This pins the current, known-imperfect behavior so a future change to the
    # function is a deliberate decision, not a silent regression.
    from app.modules.finance.chat import _normalize_markdown_tables

    ragged_with_empty_cells = (
        "| A | B |" "| --- | --- |" "| X | Y | |" "| 7 | | 9 |" "| | 11 | 12 |"
    )
    result = _normalize_markdown_tables(ragged_with_empty_cells)
    assert "7" in result and "9" in result and "11" in result and "12" in result
    assert result.count("\n") > 0  # it is reflowed, not left as one line


def test_normalize_markdown_tables_does_not_truncate_a_ragged_row():
    # A stray "|" (e.g. inside a property/vendor name pulled from the tenant's own workbook, or
    # the model emitting an uneven row) makes the pipe count not divide evenly by the column
    # width. Reflowing anyway would silently drop the leftover characters, truncating a real
    # figure with no signal -- leave the line untouched instead of guessing.
    from app.modules.finance.chat import _normalize_markdown_tables

    ragged = (
        "| Rank | Property | |------|----------|"
        "| 1 | 12 | High Street |"  # one extra "|" inside this row's second cell
    )
    assert _normalize_markdown_tables(ragged) == ragged


async def test_service_chat_reflows_a_table_the_model_glued_onto_one_line():
    # EL-1015: the model sometimes returns an entire markdown table (header, separator, every
    # data row) as one unbroken line instead of one row per line. That is not valid GFM, so the
    # frontend renders it as raw pipes instead of a table even though remark-gfm already handles
    # a properly-formatted table (EL-1013). Recover the real column count from the dashes
    # separator row (an unambiguous marker) and put each row back on its own line.
    glued = (
        "Here are the top properties:\n\n"
        "| Rank | Property | Net Profit | "
        "|------|----------|-----------|"
        "| 1 | Mandrell & Mauleverer Rd | £178,064.80 |"
        "| 2 | 1 Waterson Street | £137,838.40 |"
    )
    svc = FinanceService(provider=StubProvider(canned=glued))
    await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                     content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result = await svc.chat("acme", "agent-1", "What are the top properties by net profit?")

    assert "||" not in result["reply"]
    table_lines = [ln for ln in result["reply"].split("\n") if ln.strip().startswith("|")]
    assert table_lines == [
        "| Rank | Property | Net Profit |",
        "|------|----------|-----------|",
        "| 1 | Mandrell & Mauleverer Rd | £178,064.80 |",
        "| 2 | 1 Waterson Street | £137,838.40 |",
    ]


async def test_service_chat_without_data_does_not_hallucinate():
    svc = FinanceService(provider=StubProvider(canned="SHOULD NOT BE CALLED"))
    result = await svc.chat("acme", "agent-empty", "How much profit did we make?")
    assert "don't have any financial data" in result["reply"]
    assert result["reply"] != "SHOULD NOT BE CALLED"


async def test_service_delete_is_tenant_scoped():
    svc = FinanceService(provider=StubProvider(canned="ok"))
    stored = await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    # A different tenant cannot delete acme's file.
    other = await svc.delete_files("beta", "agent-1", [stored.id])
    assert other["deletedCount"] == 0
    mine = await svc.delete_files("acme", "agent-1", [stored.id])
    assert mine["deletedCount"] == 1


# --- router (API level) ----------------------------------------------------------------------

def _seed_api_worker(client, cookies, agent_id: str) -> None:
    import asyncio

    tenant = decode_token(cookies["qiko_access"], expected_type="access")["tenant"]
    if asyncio.run(client.app.state.agents.get(tenant, agent_id)) is None:
        asyncio.run(client.app.state.agents.store.add(
            Agent(id=agent_id, tenant=tenant, name=agent_id)))


def _upload_via_api(client, cookies, agent_id: str):
    _seed_api_worker(client, cookies, agent_id)
    return client.post(
        f"/api/v1/finance/{agent_id}/upload", cookies=cookies,
        files={"file": ("portfolio.xlsx", _build_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_finance_upload_and_analytics_via_api(client):
    cookies = auth_cookies(tenant="acme")
    up = _upload_via_api(client, cookies, "agent-fin-1")
    assert up.status_code == 201, up.text
    assert up.json()["totalRows"] == 7

    analytics = client.get("/api/v1/finance/agent-fin-1/analytics", cookies=cookies)
    assert analytics.status_code == 200
    body = analytics.json()
    assert body["summary"]["propertyCount"] == 2
    assert body["totalFiles"] == 1
    assert body["currency"] == "GBP"

    files = client.get("/api/v1/finance/agent-fin-1/files", cookies=cookies)
    assert files.json()["totalFiles"] == 1


def test_finance_analytics_month_filter_via_api(client):
    cookies = auth_cookies(tenant="acme")
    _seed_api_worker(client, cookies, "agent-fin-month")
    client.post("/api/v1/finance/agent-fin-month/upload", cookies=cookies,
                files={"file": ("p.xlsx", _two_month_workbook(),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    full = client.get("/api/v1/finance/agent-fin-month/analytics", cookies=cookies).json()
    assert full["summary"]["totalIncome"] == 1500.0

    feb = client.get("/api/v1/finance/agent-fin-month/analytics?month=2024-02", cookies=cookies).json()
    assert feb["summary"]["totalIncome"] == 1000.0
    assert feb["summary"]["totalExpenses"] == 120.0
    assert feb["dateRange"]["from"] == "2024-02-01" and feb["dateRange"]["to"] == "2024-02-29"


def test_finance_upload_requires_editor_role(client):
    viewer = auth_cookies(tenant="acme", role="viewer")
    resp = _upload_via_api(client, viewer, "agent-fin-2")
    assert resp.status_code == 403


def test_finance_rejects_non_xlsx_upload(client):
    cookies = auth_cookies(tenant="acme")
    _seed_api_worker(client, cookies, "agent-fin-3")
    resp = client.post("/api/v1/finance/agent-fin-3/upload", cookies=cookies,
                       files={"file": ("evil.xlsx", b"not a workbook",
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert resp.status_code == 422


def test_finance_is_tenant_isolated(client):
    cookies_a = auth_cookies(tenant="acme")
    cookies_b = auth_cookies(tenant="beta")
    _upload_via_api(client, cookies_a, "agent-fin-4")
    # Beta sees no files for the same agent id (BOLA-safe).
    other = client.get("/api/v1/finance/agent-fin-4/files", cookies=cookies_b)
    assert other.json()["totalFiles"] == 0
    assert other.json()["availableMonths"] == []
    other_analytics = client.get("/api/v1/finance/agent-fin-4/analytics", cookies=cookies_b)
    assert other_analytics.json()["summary"]["propertyCount"] == 0


def test_finance_chat_and_history_via_api(client):
    cookies = auth_cookies(tenant="acme")
    _upload_via_api(client, cookies, "agent-fin-5")
    chat = client.post("/api/v1/finance/agent-fin-5/chat", cookies=cookies,
                       json={"message": "Summarise total income"})
    assert chat.status_code == 200
    assert chat.json()["intent"] in ("data", "general")

    history = client.get("/api/v1/finance/agent-fin-5/history", cookies=cookies)
    assert history.status_code == 200
    assert history.json()["count"] == 2


def test_finance_delete_files_via_api(client):
    cookies = auth_cookies(tenant="acme")
    up = _upload_via_api(client, cookies, "agent-fin-6")
    namespace = up.json()["namespace"]
    file_id = namespace.rsplit("-", 1)[-1]
    resp = client.request("DELETE", "/api/v1/finance/agent-fin-6/files", cookies=cookies,
                          json={"ids": [file_id]})
    assert resp.status_code == 200
    assert resp.json()["deletedCount"] == 1


# --- Postgres-backed stores (exercised on aiosqlite, like test_db_backed_sales_intelligence) --

async def _finance_sessionmaker(path: str):
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.db.base import Base
    from app.db.models import Agent as AgentRow
    from app.db.models import Tenant

    engine = create_async_engine(f"sqlite+aiosqlite:///{path}")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    sm = async_sessionmaker(engine, expire_on_commit=False)
    async with sm() as session, session.begin():
        if await session.get(Tenant, "acme") is None:
            session.add(Tenant(id="acme", name="Acme"))
            session.add(AgentRow(id="agent-1", tenant_id="acme", name="Agent", enabled=True))
    return sm, engine


async def test_sql_finance_file_store_persists_and_is_tenant_scoped(tmp_path):
    from app.db.stores import SqlFinanceChatStore, SqlFinanceFileStore

    db_path = str(tmp_path / "finance.db")
    sm, engine = await _finance_sessionmaker(db_path)
    svc = FinanceService(files=SqlFinanceFileStore(sm), chats=SqlFinanceChatStore(sm),
                         provider=StubProvider())
    stored = await svc.upload("acme", "agent-1", _build_workbook(), filename="p.xlsx",
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    await svc.chat("acme", "agent-1", "total income?")
    await engine.dispose()  # simulate a redeploy

    sm2, _ = await _finance_sessionmaker(db_path)
    svc2 = FinanceService(files=SqlFinanceFileStore(sm2), chats=SqlFinanceChatStore(sm2),
                          provider=StubProvider())
    analytics = await svc2.analytics("acme", "agent-1")
    assert analytics["summary"]["propertyCount"] == 2  # data survived the "restart"
    # Another tenant cannot read acme's file or history (BOLA-safe).
    assert await SqlFinanceFileStore(sm2).get("beta", "agent-1", stored.id) is None
    assert (await svc2.history("beta", "agent-1"))["count"] == 0
    assert (await svc2.history("acme", "agent-1"))["count"] == 2


def test_parser_maps_description_column_to_expenditure_desc():
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append([
        "Property Reference", "Unit Ref", "Unit Description", "Description",
        "Heading", "Supname", "Nett", "Vat", "Gross", "Exp date",
    ])
    exp.append(["A001", "A001/1", "Flat 1, 12 High Street", "Leaking radiator",
                "Repairs", "Acme Maintenance", 100.0, 20.0, 120.0, "01/03/2024"])
    exp.append(["A001", "A001/1", "Flat 1, 12 High Street", "Window cleaning",
                "Cleaning", "Sparkle Clean", 50.0, 10.0, 60.0, "05/03/2024"])
    buf = io.BytesIO()
    wb.save(buf)
    expenses, _ = parse_workbook(buf.getvalue())
    assert len(expenses) == 2
    assert expenses[0].expenditure_desc == "Leaking radiator"
    assert expenses[1].expenditure_desc == "Window cleaning"
    assert expenses[0].property_desc == ""  # Description should not become property name


def test_unit_ledger_falls_back_to_income_description_when_no_expenses():
    # EL-1018: a unit with zero expense rows (Gross=0, Invoices=0) still has a description in the
    # income sheet ("Unit Description" column), but _build_unit_ledger used to source
    # unitDescription only from expense rows, so a unit with no expenses always showed a blank
    # description even though income data (and its description) was right there.
    wb = Workbook()
    exp = wb.active
    exp.title = "Exp YTD"
    exp.append(_EXPENSE_HEADER)
    exp.append(["A001", "12 High Street", "A001/1", "Flat 1, 12 High Street", "Repairs",
                "Acme Maintenance", 100.0, 20.0, 120.0, "01/03/2024"])
    # A001/2 has no expense row at all -> Gross=0, Invoices=0 for that unit.

    inc = wb.create_sheet("Income Receivable YTD")
    inc.append(_INCOME_HEADER)
    inc.append(["A001", "A001/1", "Flat 1, 12 High Street", 500.0, 0.0, "05/03/2024"])
    inc.append(["A001", "A001/2", "Flat 2, 12 High Street", 400.0, 0.0, "05/03/2024"])

    buf = io.BytesIO()
    wb.save(buf)
    expenses, incomes = parse_workbook(buf.getvalue())
    result = compute_analytics(expenses, incomes)

    prop_a = next(p for p in result["properties"] if p["propertyRef"] == "A001")
    unit2 = next(u for u in prop_a["unitLedger"] if u["unitRef"] == "A001/2")
    assert unit2["totalGross"] == 0.0
    assert unit2["invoiceCount"] == 0
    assert unit2["unitDescription"] == "Flat 2, 12 High Street"


def test_analytics_enriches_property_ref_and_monthly_trends():
    expenses, incomes = parse_workbook(_build_workbook())
    result = compute_analytics(expenses, incomes)

    prop_a = next(p for p in result["properties"] if p["propertyRef"] == "A001")
    assert prop_a["name"] == "12 High Street"
    assert "monthlyTrends" in prop_a
    assert any(t["month"] == "2024-03" for t in prop_a["monthlyTrends"])
    assert prop_a["vendorCount"] == 2
    assert prop_a["unitCount"] == 1

    assert "monthlyTrends" in result["charts"]

    unit = prop_a["unitLedger"][0]
    assert unit["propertyRef"] == "A001"
    assert unit["propertyName"] == "12 High Street"
    assert unit["expenses"][0]["propertyRef"] == "A001"
    assert unit["expenses"][0]["unitRef"] == "A001/1"
    assert unit["expenses"][0]["expenditureDescription"] != ""

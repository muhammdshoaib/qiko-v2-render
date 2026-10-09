import io
import json

import pytest

from app.modules.knowledge.parsing import (
    UnsupportedDocumentError,
    extract_text,
    normalize_text,
)


def test_normalize_collapses_whitespace():
    assert normalize_text("a  \t b\r\n\n\n\nc  ") == "a b\n\nc"


def test_txt_and_md():
    assert extract_text("notes.txt", b"hello world") == "hello world"
    assert extract_text("r.md", b"# Title\n\ntext") == "# Title\n\ntext"


def test_csv():
    out = extract_text("d.csv", b"name,role\nAda,owner\nBob,viewer\n")
    assert "name, role" in out and "Ada, owner" in out


def test_json_and_jsonl():
    out = extract_text("d.json", json.dumps({"company": "IDG", "certs": ["ISO 9001"]}).encode())
    assert "company: IDG" in out and "ISO 9001" in out
    jl = b'{"q": "a"}\n{"q": "b"}\n'
    assert "q: a" in extract_text("d.jsonl", jl) and "q: b" in extract_text("d.jsonl", jl)


def test_docx_real_extraction():
    from docx import Document

    doc = Document()
    doc.add_paragraph("IDG holds ISO 9001 certification.")
    buf = io.BytesIO()
    doc.save(buf)
    assert "ISO 9001 certification" in extract_text("bid.docx", buf.getvalue())


def test_docx_extracts_table_content():
    # Regression: the docx reader used to read only paragraphs, silently dropping tables -- and in
    # tender packs (Annex templates / compliance matrices) the substance lives in the tables.
    from docx import Document

    doc = Document()
    doc.add_paragraph("Technical Proposal Template")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Requirement"
    table.cell(0, 1).text = "Response"
    table.cell(1, 0).text = "Security clearance"
    table.cell(1, 1).text = "SC-cleared personnel provided"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("annex.docx", buf.getvalue())
    assert "Security clearance" in text
    assert "SC-cleared personnel provided" in text


def test_xlsx_real_extraction():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Certs"
    ws.append(["standard", "valid_until"])
    ws.append(["ISO 9001", 2027])
    buf = io.BytesIO()
    wb.save(buf)
    out = extract_text("certs.xlsx", buf.getvalue())
    assert "# Certs" in out and "ISO 9001, 2027" in out


def test_unsupported_type_rejected():
    with pytest.raises(UnsupportedDocumentError):
        extract_text("malware.exe", b"MZ")


def test_html_strips_tags_and_keeps_visible_text():
    html = b"""
    <html><head><title>ignored</title><style>.x{color:red}</style></head>
    <body>
      <h1>Tender BOC 2171</h1>
      <p>Provision of Static Security Services.</p>
      <script>var secret = "should not appear";</script>
      <ul><li>Deadline: 2026-08-01</li><li>Location: Iraq</li></ul>
    </body></html>
    """
    out = extract_text("page.html", html)
    assert "Tender BOC 2171" in out
    assert "Provision of Static Security Services." in out
    assert "Deadline: 2026-08-01" in out
    assert "should not appear" not in out
    assert "color:red" not in out


def test_htm_extension_also_supported():
    assert "hello" in extract_text("page.htm", b"<p>hello</p>")


def _deck(builder):
    from pptx import Presentation

    prs = Presentation()
    builder(prs)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_real_extraction():
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Incident Response"
        slide.placeholders[1].text = "Control room alerted within five minutes."

    out = extract_text("conops.pptx", _deck(build))
    assert "[slide 1]" in out
    assert "Incident Response" in out
    assert "Control room alerted within five minutes." in out


def test_pptx_labels_each_slide_so_citations_can_point_at_one():
    def build(prs):
        for title in ("Armed guarding", "Unarmed guarding"):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = title

    out = extract_text("deck.pptx", _deck(build))
    assert "[slide 1]" in out and "[slide 2]" in out
    assert out.index("Armed guarding") < out.index("[slide 2]")


def test_pptx_reads_tables():
    from pptx.util import Inches

    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Role"
        table.cell(0, 1).text = "Headcount"
        table.cell(1, 0).text = "Static guard"
        table.cell(1, 1).text = "250"

    out = extract_text("mobplan.pptx", _deck(build))
    assert "Role | Headcount" in out
    assert "Static guard | 250" in out


def test_pptx_reads_grouped_shapes():
    # Org charts are almost always one grouped shape; a non-recursive walk returns nothing at all.
    from pptx.util import Inches

    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shapes = slide.shapes
        first = shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        first.text_frame.text = "Country Manager"
        second = shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
        second.text_frame.text = "Deputy Country Manager"
        shapes._spTree.remove(first._element)
        shapes._spTree.remove(second._element)
        group = shapes.add_group_shape()
        group.shapes._spTree.append(first._element)
        group.shapes._spTree.append(second._element)

    out = extract_text("orgchart.pptx", _deck(build))
    assert "Country Manager" in out
    assert "Deputy Country Manager" in out


def test_pptx_reads_speaker_notes():
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Mobilisation"
        slide.notes_slide.notes_text_frame.text = "Thirty day mobilisation across twelve sites."

    out = extract_text("mob.pptx", _deck(build))
    assert "Thirty day mobilisation across twelve sites." in out


def test_legacy_xls_is_rejected_not_mis_parsed():
    # openpyxl cannot read the legacy BIFF .xls format, so mapping "xls" at it turned a bad upload
    # into a 500 instead of the 400 every other unsupported type gets.
    with pytest.raises(UnsupportedDocumentError):
        extract_text("mobplan.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")


def test_supported_cannot_drift_from_the_extractors():
    from app.modules.knowledge.parsing import SUPPORTED, _EXTRACTORS

    assert SUPPORTED == set(_EXTRACTORS)

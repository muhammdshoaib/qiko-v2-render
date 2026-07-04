"""Downloadable file generation. An agent produces real office artifacts (DOCX/XLSX/CSV/PDF) from
structured content and stores them in S3 (tenant-scoped), returning a download key. This is a tool
the agent calls, not a free-form model output. (PPTX and chart-image PNGs are a follow-up needing
python-pptx / matplotlib.)
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from app.modules.storage.backend import InMemoryStorage, StorageBackend
from app.modules.storage.keys import object_key

SUPPORTED_FORMATS = {"csv", "xlsx", "docx", "pdf"}


class UnsupportedFormatError(Exception):
    pass


def build_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8")


def build_xlsx(sheet: str, rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet or "Sheet1"
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_docx(title: str, paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(title: str, paragraphs: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _, height = A4
    y = height - 72
    if title:
        c.setFont("Helvetica-Bold", 16)
        c.drawString(72, y, title)
        y -= 28
    c.setFont("Helvetica", 11)
    for p in paragraphs:
        c.drawString(72, y, p[:110])
        y -= 16
    c.showPage()
    c.save()
    return buf.getvalue()


@dataclass
class FileGenerator:
    storage: StorageBackend = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.storage is None:
            self.storage = InMemoryStorage()

    def build(self, fmt: str, spec: dict) -> bytes:
        if fmt == "csv":
            return build_csv(spec.get("rows", []))
        if fmt == "xlsx":
            return build_xlsx(spec.get("sheet", "Sheet1"), spec.get("rows", []))
        if fmt == "docx":
            return build_docx(spec.get("title", ""), spec.get("paragraphs", []))
        if fmt == "pdf":
            return build_pdf(spec.get("title", ""), spec.get("paragraphs", []))
        raise UnsupportedFormatError(f"unsupported format: {fmt}")

    def generate(self, tenant: str, fmt: str, filename: str, spec: dict) -> str:
        """Build the artifact and store it under the tenant's exports namespace; returns the key."""
        if fmt not in SUPPORTED_FORMATS:
            raise UnsupportedFormatError(f"unsupported format: {fmt}")
        data = self.build(fmt, spec)
        key = object_key(tenant, "exports", f"{filename}.{fmt}")
        self.storage.put(tenant, key, data, content_type="application/octet-stream")
        return key

"""Excel ingest + NL search, ported from the mcp-parser as a knowledge sub-feature.

A workbook is parsed into per-row records (header: value), embedded (Titan V2 in prod) and stored
in pgvector per tenant via KnowledgeService, then searchable with NL queries. Exposed as
schema-validated agentcore tools. The legacy remote-eval RCE is NOT ported.
"""
from __future__ import annotations

import io
from dataclasses import dataclass

from pydantic import BaseModel, Field

from app.modules.agentcore.tools import ToolRegistry
from app.modules.knowledge.service import KnowledgeService


def workbook_to_records(data: bytes) -> list[str]:
    """Each data row becomes a 'col: value, ...' record line (header-aware)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    records: list[str] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        for row in rows[1:]:
            pairs = [f"{header[i]}: {v}" for i, v in enumerate(row) if v is not None and i < len(header)]
            if pairs:
                records.append(f"[{ws.title}] " + ", ".join(pairs))
    return records


@dataclass
class ExcelService:
    knowledge: KnowledgeService

    async def store_workbook(self, tenant: str, source_name: str, data: bytes) -> int:
        records = workbook_to_records(data)
        for i, rec in enumerate(records):
            await self.knowledge.ingest(tenant, f"excel-{source_name}-{i}", rec)
        return len(records)

    async def search(self, tenant: str, query: str, *, k: int = 5) -> list[str]:
        return [p.text for p in await self.knowledge.retrieve(tenant, query, k=k)]


class StoreArgs(BaseModel):
    source_name: str = Field(min_length=1)


class SearchArgs(BaseModel):
    query: str = Field(min_length=1)
    k: int = Field(default=5, ge=1, le=50)


def build_excel_tools(service: ExcelService, tenant: str, workbook: bytes) -> ToolRegistry:
    reg = ToolRegistry()

    @reg.register("store_excel_data", StoreArgs)
    async def _store(a: StoreArgs) -> int:
        return await service.store_workbook(tenant, a.source_name, workbook)

    @reg.register("search_data", SearchArgs)
    async def _search(a: SearchArgs) -> list[str]:
        return await service.search(tenant, a.query, k=a.k)

    return reg

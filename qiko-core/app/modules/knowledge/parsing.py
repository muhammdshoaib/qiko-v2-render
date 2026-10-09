"""Multi-format document text extraction for knowledge ingestion.

Turns an uploaded file (PDF/DOCX/PPTX/CSV/XLSX/JSON/JSONL/TXT/MD) into normalised plain text ready
for chunk -> embed. Heavy binary readers are isolated per format so the ingestion pipeline stays
format-agnostic. Unsupported types are rejected (no silent empty ingest).
"""
from __future__ import annotations

import csv
import io
import json
import re

# Hard per-upload cap for document endpoints that parse an uploaded file (rfp /extract, knowledge
# /ingest-file). Mirrors finance.service.MAX_UPLOAD_BYTES and the global request-size middleware --
# defence in depth so a huge upload can't exhaust memory before parsing.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB


class UnsupportedDocumentError(Exception):
    pass


def normalize_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_text(data: bytes) -> str:
    return data.decode("utf-8", "ignore")


def _read_csv(data: bytes) -> str:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8", "ignore"))))
    return "\n".join(", ".join(cell for cell in row) for row in rows if any(row))


def _json_to_text(obj, prefix: str = "") -> str:
    lines: list[str] = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            lines.append(_json_to_text(v, f"{prefix}{k}: "))
    elif isinstance(obj, list):
        for item in obj:
            lines.append(_json_to_text(item, prefix))
    else:
        lines.append(f"{prefix}{obj}")
    return "\n".join(x for x in lines if x)


def _read_json(data: bytes) -> str:
    return _json_to_text(json.loads(data.decode("utf-8", "ignore")))


def _read_jsonl(data: bytes) -> str:
    out = []
    for line in data.decode("utf-8", "ignore").splitlines():
        line = line.strip()
        if line:
            out.append(_json_to_text(json.loads(line)))
    return "\n".join(out)


def _read_docx(data: bytes) -> str:
    # Walk the document body in order so we capture BOTH paragraphs AND tables -- in tender packs
    # (e.g. an Annex "Technical Proposal Template" or a compliance matrix) the substance lives in
    # the tables, and a paragraphs-only read silently drops the bulk of the content. Also pull in
    # header/footer text (deduped, since it repeats per section).
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    parts: list[str] = []

    def _table_lines(tbl: Table) -> None:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]  # version-robust: 'p' | 'tbl' | ...
        if tag == "p":
            text = Paragraph(child, doc).text
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            _table_lines(Table(child, doc))

    seen: set[str] = set()
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in hf.paragraphs:
                t = p.text.strip()
                if t and t not in seen:
                    seen.add(t)
                    parts.append(t)

    return "\n".join(parts)


def _read_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(", ".join(cells))
    return "\n".join(parts)


def _read_pptx(data: bytes) -> str:
    # Slide decks carry the operational detail in tender packs: org charts, CONOPs, incident-response
    # and mobilisation flows. Three things make a naive read return almost nothing. The text sits in
    # free-floating shapes rather than a document body; an org chart is nearly always a single GROUP
    # shape, so a non-recursive walk sees one empty container; and the prose a bid writer actually
    # wants is frequently in the speaker notes rather than on the slide.
    #
    # Slides are labelled so a citation can point at one, and so the chunker has a real boundary to
    # split on instead of fusing a whole deck into one window.
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []

    def _walk(shape) -> None:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                _walk(child)
            return
        if shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
            return
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                parts.append(text)

    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"[slide {index}]")
        for shape in slide.shapes:
            _walk(shape)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[slide {index} notes]")
                parts.append(notes)
    return "\n".join(parts)


def _read_pdf(data: bytes) -> str:  # pragma: no cover - exercised via integration (binary PDF)
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


class _HtmlTextExtractor:
    """Minimal HTML -> text (stdlib html.parser, no new dependency): drops <script>/<style>,
    keeps visible text, and inserts a newline at block-level tag boundaries so paragraphs and
    headings don't run together -- good enough for "paste a tender webpage URL", which just
    needs readable text fed into the same chunk/extract pipeline as an uploaded document."""
    _BLOCK_TAGS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}
    _SKIP_TAGS = {"script", "style", "noscript", "head"}

    def __init__(self) -> None:
        from html.parser import HTMLParser

        self._chunks: list[str] = []
        self._skip_depth = 0

        class _Parser(HTMLParser):
            def handle_starttag(p_self, tag, attrs):  # noqa: N805 - HTMLParser convention
                if tag in self._SKIP_TAGS:
                    self._skip_depth += 1
                elif tag in self._BLOCK_TAGS:
                    self._chunks.append("\n")

            def handle_endtag(p_self, tag):  # noqa: N805
                if tag in self._SKIP_TAGS:
                    self._skip_depth = max(0, self._skip_depth - 1)
                elif tag in self._BLOCK_TAGS:
                    self._chunks.append("\n")

            def handle_data(p_self, data):  # noqa: N805
                if self._skip_depth == 0 and data.strip():
                    self._chunks.append(data)

        self._parser = _Parser()

    def extract(self, html: str) -> str:
        self._parser.feed(html)
        self._parser.close()
        return "".join(self._chunks)


def _read_html(data: bytes) -> str:
    return _HtmlTextExtractor().extract(data.decode("utf-8", "ignore"))


# "xls" is deliberately absent. It was mapped at _read_xlsx, but openpyxl reads only the OOXML
# format and raises on a legacy BIFF .xls, so claiming support turned a bad upload into a 500
# instead of the 400 every other unsupported type gets. Adding it back means a real BIFF reader.
_EXTRACTORS = {
    "txt": _read_text, "md": _read_text, "csv": _read_csv,
    "json": _read_json, "jsonl": _read_jsonl,
    "docx": _read_docx, "xlsx": _read_xlsx, "pptx": _read_pptx, "pdf": _read_pdf,
    "html": _read_html, "htm": _read_html,
}

# Derived, never hand-listed: the two drifted apart once already and a stale allowlist either
# rejects a format the parser handles or promises one it does not.
SUPPORTED = frozenset(_EXTRACTORS)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in _EXTRACTORS:
        raise UnsupportedDocumentError(f"unsupported document type: .{ext or '?'}")
    return normalize_text(_EXTRACTORS[ext](data))
